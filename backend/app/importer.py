"""Import of contract lists (CSV/Excel), e.g. from an own spreadsheet.

Files are only parsed in memory and never stored.
"""

import csv
import io
import re
from datetime import date, datetime

from .categories import CATEGORIES

# --------------------------------------------------------------------------- reading tables


class ImportError_(Exception):
    pass


def _decode(blob: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return blob.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportError_("Datei konnte nicht gelesen werden")


def read_rows(blob: bytes, filename: str) -> list[list[str]]:
    """All rows of a CSV or Excel file as lists of strings."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or blob[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ImportError_("Excel-Datei konnte nicht gelesen werden") from exc
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(
                [
                    "" if v is None else (v.strftime("%d.%m.%Y") if isinstance(v, datetime | date) else str(v))
                    for v in row
                ]
            )
        return rows
    text = _decode(blob)
    sample = text[:5000]
    delimiter = max([";", ",", "\t"], key=sample.count)
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def parse_amount(value) -> float | None:
    if value is None:
        return None
    s = str(value).strip().replace("€", "").replace("EUR", "").replace(" ", "").replace(" ", "")
    if not s:
        return None
    negative = s.endswith("-") or s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = s.strip("-+()")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        v = float(s)
    except ValueError:
        return None
    return -v if negative else v


def parse_date(value) -> date | None:
    s = str(value or "").strip()[:10]
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _find_header(rows: list[list[str]], keywords: dict[str, tuple[str, ...]], required: tuple[str, ...]):
    """Locate the header row (exports often start with some info lines) and map columns.

    For every field the keywords are tried in order of preference, so "Betrag" wins over
    a column like "Umsatzart"; an exact label wins over one that only contains the keyword
    (Sparkasse: "Betrag" vs. "Lastschrift Ursprungsbetrag").
    """
    matchers = (
        lambda w, label: label == w,
        lambda w, label: label.startswith(w),
        lambda w, label: w in label,
    )

    for idx, row in enumerate(rows[:40]):
        labels = [cell.lower().strip() for cell in row]
        cols: dict[str, int] = {}
        used: set[int] = set()
        for key, words in keywords.items():
            for w in words:
                j = next(
                    (j for match in matchers for j, label in enumerate(labels) if j not in used and match(w, label)),
                    None,
                )
                if j is not None:
                    cols[key] = j
                    used.add(j)
                    break
        if all(k in cols for k in required) and len(cols) >= 2:
            return idx, cols
    return None, {}


# --------------------------------------------------------------------------- categories

CATEGORY_HINTS = [
    ("electricity", ("strom", "stadtwerke", "energie", "eon", "e.on", "vattenfall", "enbw", "lichtblick", "tibber")),
    ("gas", ("erdgas", "gasversorgung")),
    ("mobile", ("mobilfunk", "congstar", "o2 ", "telefonica", "freenet", "aldi talk", "drillisch", "mobile")),
    ("internet", ("dsl", "glasfaser", "kabel", "1&1", "vodafone", "telekom", "pyur", "netcologne", "m-net")),
    ("car_insurance", ("kfz", "huk", "allianz", "devk", "kraftfahrt", "autoversicherung")),
    ("legal_protection", ("rechtsschutz", "arag", "roland", "advocard")),
    ("liability", ("haftpflicht",)),
    ("household", ("hausrat",)),
    ("insurance", ("versicherung", "ergo", "axa", "generali", "debeka", "barmenia", "provinzial")),
    ("streaming", ("netflix", "spotify", "disney", "dazn", "sky ", "amazon prime", "apple.com", "youtube", "audible")),
    ("fitness", ("fitness", "fitx", "mcfit", "urban sports", "sportnavi", "wellpass", "verein", "gym")),
]


def guess_category(text: str) -> str:
    t = f" {text.lower()} "
    for key, words in CATEGORY_HINTS:
        if any(w in t for w in words):
            return key
    return "other"


# --------------------------------------------------------------------------- contract lists

LIST_COLUMNS = {
    "name": ("bezeichnung", "vertrag", "tarif", "name", "titel"),
    "provider": ("anbieter", "provider", "firma", "gesellschaft", "versicherer"),
    "category": ("kategorie", "art", "typ", "category", "bereich"),
    "amount": ("preis", "betrag", "kosten", "beitrag", "amount"),
    "interval": ("intervall", "zahlweise", "rhythmus", "turnus", "interval"),
    "start_date": ("beginn", "start", "abschluss", "seit"),
    "customer_number": ("kundennummer", "kunden-nr", "kd-nr"),
    "contract_number": ("vertragsnummer", "vertrags-nr", "police", "versicherungsnummer"),
    "min_term_months": ("laufzeit", "mindestlaufzeit"),
    "notice_value": ("kündigungsfrist", "kuendigungsfrist", "frist"),
    "notes": ("notiz", "bemerkung", "kommentar", "notes"),
}

INTERVAL_WORDS = {
    "monthly": ("monat", "mtl", "monthly"),
    "quarterly": ("viertel", "quartal", "quarter"),
    "half_yearly": ("halbj", "half"),
    "yearly": ("jahr", "jährl", "jaehrl", "p.a", "yearly", "annual"),
}


def _category_key(text: str, fallback_text: str) -> str:
    t = (text or "").strip().lower()
    for c in CATEGORIES:
        if t and (t == c["key"] or t in c["label"].lower() or c["label"].lower() in t):
            return c["key"]
    return guess_category(f"{text} {fallback_text}")


def _interval_key(text: str) -> str:
    t = (text or "").lower()
    for key, words in INTERVAL_WORDS.items():
        if any(w in t for w in words):
            return key
    return "monthly"


def _int(text) -> int | None:
    m = re.search(r"\d+", str(text or ""))
    return int(m.group()) if m else None


def read_contract_list(blob: bytes, filename: str) -> list[dict]:
    rows = read_rows(blob, filename)
    header_idx, cols = _find_header(rows, LIST_COLUMNS, ("amount",))
    if header_idx is None or not ({"name", "provider"} & cols.keys()):
        raise ImportError_("Spalten nicht erkannt – benötigt: Bezeichnung oder Anbieter und Preis. Nutze die Vorlage.")
    out = []
    for row in rows[header_idx + 1 :]:

        def cell(key, row=row):
            j = cols.get(key)
            return row[j].strip() if j is not None and j < len(row) and row[j] else ""

        name, provider = cell("name"), cell("provider")
        amount = parse_amount(cell("amount"))
        if not (name or provider) or amount is None:
            continue
        notice = cell("notice_value")
        out.append(
            {
                "name": name or provider,
                "provider": provider or None,
                "category": _category_key(cell("category"), f"{name} {provider}"),
                "amount": abs(amount),
                "interval": _interval_key(cell("interval")),
                "start_date": parse_date(cell("start_date")) or date.today().replace(day=1),
                "customer_number": cell("customer_number") or None,
                "contract_number": cell("contract_number") or None,
                "min_term_months": _int(cell("min_term_months")),
                "notice_value": _int(notice),
                "notice_unit": "weeks"
                if "woche" in notice.lower()
                else "days"
                if "tag" in notice.lower()
                else "months",
                "notes": cell("notes") or None,
            }
        )
    if not out:
        raise ImportError_("Keine Verträge in der Datei gefunden")
    return out


TEMPLATE = (
    "Bezeichnung;Anbieter;Kategorie;Preis;Zahlweise;Beginn;Kundennummer;Vertragsnummer;Laufzeit (Monate);"
    "Kündigungsfrist;Notizen\n"
    "MagentaMobil M;Telekom;Mobilfunk;42,95;monatlich;01.04.2024;123456;;24;1 Monat;\n"
    "Kfz-Versicherung;HUK-COBURG;Kfz-Versicherung;534,00;jährlich;01.01.2024;;KFZ-778;12;1 Monat;Golf\n"
)
